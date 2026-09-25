import os
import sys
import io
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import telebot
from telebot import types

TOKEN = '8574451645:AAER4vkfOzip0KHolGmXqaeKfdmtN0f9bdk'

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# 1. Dynamic, bulletproof Excel file detection
def find_excel_file():
    # Priority candidates
    candidates = [
        os.path.join(BASE_DIR, 'konkur_data.xlsx'),
        os.path.join(BASE_DIR, 'آخرین رتبه های قبولی تجربی 1400 تا 1403 - نسخه نهایی.xlsx'),
        'konkur_data.xlsx',
        'آخرین رتبه های قبولی تجربی 1400 تا 1403 - نسخه نهایی.xlsx'
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
            
    # Search all .xlsx files in BASE_DIR and current directory
    for search_dir in [BASE_DIR, os.getcwd()]:
        if os.path.exists(search_dir):
            for f in os.listdir(search_dir):
                if f.endswith('.xlsx') and not f.startswith('~$'):
                    return os.path.join(search_dir, f)
    return None

EXCEL_PATH = find_excel_file()
if not EXCEL_PATH or not os.path.exists(EXCEL_PATH):
    print("❌ خطا: فایل اکسل پایگاه داده در پوشه ربات پیدا نشد!")
    print(f"مسیر جستجو: {BASE_DIR}")
    sys.exit(1)

print(f"📁 فایل پایگاه داده شناسایی شد: {os.path.basename(EXCEL_PATH)}")
print("در حال بارگذاری پایگاه داده جامع کنکور تجربی...")

# Resilient sheet reading
xl = pd.ExcelFile(EXCEL_PATH)
sheet_names = xl.sheet_names
target_sheet = sheet_names[0]
for s in sheet_names:
    if 'تجمیع' in s or 'همه' in s:
        target_sheet = s
        break

print(f"📄 در حال خواندن شیت: {target_sheet}")
df = pd.read_excel(EXCEL_PATH, sheet_name=target_sheet)
df['رتبه در سهمیه'] = pd.to_numeric(df['رتبه در سهمیه'], errors='coerce')
df = df.dropna(subset=['رتبه در سهمیه'])
df['سهمیه'] = pd.to_numeric(df['سهمیه'], errors='coerce').fillna(0).astype(int)
print(f"✅ تعداد {len(df):,} ردیف داده با موفقیت در حافظه بارگذاری شد.")

bot = telebot.TeleBot(TOKEN)

# 2. Target Majors List in Exact Requested Order
TARGET_MAJORS = [
    'پزشکی',
    'دندانپزشکی',
    'داروسازی',
    'فیزیوتراپی',
    'دامپزشکی',
    'شنواییسنجی',
    'گفتاردرمانی',
    'اعضای مصنوعی',
    'بیناییسنجی',
    'پرستاری',
    'مامایی',
    'اتاق عمل',
    'هوشبری',
    'رادیولوژی',
    'پرتودرمانی',
    'پزشکی هستهای',
    'علوم آزمایشگاهی',
    'علوم تغذیه',
    'فوریتهای پزشکی',
    'فناوری اطلاعات سلامت',
    'کتابداری پزشکی',
    'بهداشت حرفهای',
    'بهداشت محیط',
    'بهداشت عمومی',
    'زیست سلولی-مولکولی',
    'علوم دامی'
]

def get_major_priority(major_str):
    m = str(major_str).strip()
    m_clean = m.replace('ي', 'ی').replace('ك', 'ک').replace('\u200c', '').replace(' ', '')
    
    if 'دندانپزشکی' in m_clean or 'دندان' in m_clean:
        return 1
    if 'دامپزشکی' in m_clean or 'دامپزشک' in m_clean:
        return 4
    if 'پزشکیهسته' in m_clean or 'هستهای' in m_clean or 'هسته' in m_clean:
        return 15
    if 'فوریت' in m_clean:
        return 18
    if 'کتابداری' in m_clean:
        return 20
    if 'اعضایمصنوعی' in m_clean or 'ارتوز' in m_clean or 'پروتز' in m_clean:
        return 7
    if m_clean == 'پزشکی' or m_clean.startswith('پزشکی'):
        return 0
    if 'فیزیوتراپی' in m_clean:
        return 3
    if 'شنوایی' in m_clean:
        return 5
    if 'گفتاردرمانی' in m_clean or 'گفتار' in m_clean:
        return 6
    if 'بینایی' in m_clean:
        return 8
    if 'پرستاری' in m_clean:
        return 9
    if 'مامایی' in m_clean:
        return 10
    if 'اتاقعمل' in m_clean:
        return 11
    if 'هوشبری' in m_clean:
        return 12
    if 'رادیولوژی' in m_clean or 'پرتوشناسی' in m_clean:
        return 13
    if 'پرتودرمانی' in m_clean:
        return 14
    if 'آزمایشگاهی' in m_clean:
        return 16
    if 'تغذیه' in m_clean:
        return 17
    if 'اطلاعاتسلامت' in m_clean or 'مدارکپزشکی' in m_clean:
        return 19
    if 'بهداشتحرفه' in m_clean or 'ایمنیکار' in m_clean:
        return 21
    if 'بهداشتمحیط' in m_clean:
        return 22
    if 'بهداشتعمومی' in m_clean:
        return 23
    if 'سلولی' in m_clean or 'سلولیمولکولی' in m_clean:
        return 24
    if 'علومدامی' in m_clean or 'دامپروری' in m_clean:
        return 25

    return 99

user_state = {}

def get_region_keyboard():
    markup = types.InlineKeyboardMarkup(row_width=3)
    markup.add(
        types.InlineKeyboardButton("منطقه ۱", callback_data="reg_1"),
        types.InlineKeyboardButton("منطقه ۲", callback_data="reg_2"),
        types.InlineKeyboardButton("منطقه ۳", callback_data="reg_3")
    )
    return markup

def get_source_keyboard():
    markup = types.InlineKeyboardMarkup(row_width=1)
    markup.add(
        types.InlineKeyboardButton("🌐 همه با هم (پایگاه تجمیعی کامل)", callback_data="src_all"),
        types.InlineKeyboardButton("📋 کارنامه‌های قبولی (قلم‌چی و گزینه ۲)", callback_data="src_kanoon"),
        types.InlineKeyboardButton("📊 مرجع جامع کشوری (سنجش و دانشگاه‌ها)", callback_data="src_ref")
    )
    return markup

def get_majors_keyboard():
    markup = types.InlineKeyboardMarkup(row_width=2)
    # First row: All majors
    markup.add(types.InlineKeyboardButton("🔍 همه رشته‌ها (جامع)", callback_data="maj_all"))
    
    # 26 target majors in pairs
    buttons = []
    for idx, name in enumerate(TARGET_MAJORS):
        buttons.append(types.InlineKeyboardButton(name, callback_data=f"maj_{idx}"))
    
    # Add in pairs of 2
    for i in range(0, len(buttons), 2):
        if i + 1 < len(buttons):
            markup.add(buttons[i], buttons[i+1])
        else:
            markup.add(buttons[i])
            
    return markup

# Handlers
@bot.message_handler(commands=['start', 'help'])
def send_welcome(message):
    chat_id = message.chat.id
    user_state[chat_id] = {}
    
    welcome_text = (
        "👋 **به ربات انتخاب رشته و تخمین شانس قبولی کنکور تجربی خوش آمدید!**\n\n"
        "✨ **ویژگی‌های نسخه جدید:**\n"
        "▫️ بازه تحلیلی: **۳۰٪ خوش‌بینانه** تا **۲۵٪ بدبینانه (حاشیه امن)**\n"
        "▫️ امکان تفکیک نتایج بر اساس **هر دو پایگاه داده** یا **تجمیعی**\n"
        "▫️ خروجی اکسل اختصاصی با ترتیب استاندارد رشته‌ها و طراحی شکیل\n\n"
        "📍 **لطفاً سهمیه منطقه خود را انتخاب فرمایید:**"
    )
    bot.send_message(chat_id, welcome_text, parse_mode='Markdown', reply_markup=get_region_keyboard())

@bot.callback_query_handler(func=lambda call: call.data.startswith('reg_'))
def callback_region(call):
    chat_id = call.message.chat.id
    reg_num = int(call.data.split('_')[1])
    if chat_id not in user_state:
        user_state[chat_id] = {}
    user_state[chat_id]['region'] = reg_num
    
    bot.answer_callback_query(call.id, f"منطقه {reg_num} انتخاب شد.")
    bot.edit_message_text(
        f"✅ سهمیه **منطقه {reg_num}** ثبت شد.\n\n"
        "🎯 لطفاً **رتبه در سهمیه** خود را به‌صورت عدد تایپ و ارسال فرمایید (مثلاً: `6500`):",
        chat_id=chat_id,
        message_id=call.message.message_id,
        parse_mode='Markdown'
    )
    bot.register_next_step_handler(call.message, process_rank_step)

def process_rank_step(message):
    chat_id = message.chat.id
    text = message.text.strip().replace(',', '')
    
    persian_digits = '۰۱۲۳۴۵۶۷۸۹'
    for i, p in enumerate(persian_digits):
        text = text.replace(p, str(i))
        
    if not text.isdigit():
        bot.send_message(chat_id, "⚠️ لطفاً فقط مقدار عددی رتبه را وارد فرمایید (مثلاً: 6500):")
        return bot.register_next_step_handler(message, process_rank_step)
        
    rank = int(text)
    if rank <= 0 or rank > 300000:
        bot.send_message(chat_id, "⚠️ رتبه وارد شده معتبر نیست. لطفاً مجدداً رتبه را وارد کنید:")
        return bot.register_next_step_handler(message, process_rank_step)
        
    if chat_id not in user_state:
        user_state[chat_id] = {'region': 2}
    user_state[chat_id]['rank'] = rank
    
    prompt_text = (
        f"✅ رتبه **{rank:,}** (منطقه {user_state[chat_id].get('region', 2)}) ثبت گردید.\n\n"
        "🔍 **مایلید استعلام قبولی‌ها بر اساس کدام پایگاه داده انجام شود؟**"
    )
    bot.send_message(chat_id, prompt_text, parse_mode='Markdown', reply_markup=get_source_keyboard())

@bot.callback_query_handler(func=lambda call: call.data.startswith('src_'))
def callback_source(call):
    chat_id = call.message.chat.id
    src_key = call.data.split('_')[1]
    
    source_map = {
        'all': ('همه', '🌐 همه با هم (پایگاه تجمیعی کامل)'),
        'kanoon': ('کارنامه', '📋 کارنامه‌های قبولی (قلم‌چی و گزینه ۲)'),
        'ref': ('مرجع جامع', '📊 مرجع جامع کشوری (سنجش و دانشگاه‌ها)')
    }
    src_code, src_label = source_map.get(src_key, ('همه', '🌐 همه با هم'))
    
    if chat_id not in user_state:
        user_state[chat_id] = {}
    user_state[chat_id]['source'] = src_code
    user_state[chat_id]['source_label'] = src_label
    
    bot.answer_callback_query(call.id, "منبع انتخاب شد.")
    
    bot.edit_message_text(
        f"✅ منبع استعلام: **{src_label}**\n\n"
        "📚 **اکنون رشته مدنظر خود را انتخاب فرمایید:**\n"
        "(یا نام رشته دلخواه خود را در چت تایپ فرمایید)",
        chat_id=chat_id,
        message_id=call.message.message_id,
        parse_mode='Markdown',
        reply_markup=get_majors_keyboard()
    )

@bot.callback_query_handler(func=lambda call: call.data.startswith('maj_'))
def callback_major(call):
    chat_id = call.message.chat.id
    maj_key = call.data.split('_')[1]
    
    if maj_key == 'all':
        major_name = 'همه رشته‌ها'
    else:
        idx = int(maj_key)
        major_name = TARGET_MAJORS[idx]
        
    bot.answer_callback_query(call.id, f"رشته: {major_name}")
    execute_search_and_send(chat_id, major_name, call.message.message_id)

@bot.message_handler(func=lambda msg: True)
def text_input_fallback(message):
    chat_id = message.chat.id
    if chat_id in user_state and 'rank' in user_state[chat_id]:
        major_name = message.text.strip()
        execute_search_and_send(chat_id, major_name)
    else:
        bot.send_message(chat_id, "💡 برای شروع استعلام انتخاب رشته، دستور /start را ارسال فرمایید.")

def execute_search_and_send(chat_id, major_name, edit_msg_id=None):
    state = user_state.get(chat_id, {})
    rank = state.get('rank', 5000)
    region = state.get('region', 2)
    src_filter = state.get('source', 'همه')
    src_label = state.get('source_label', '🌐 همه با هم')
    
    # 30% optimistic (better rank), 25% pessimistic (worse rank)
    min_rank = int(rank * 0.70)
    max_rank = int(rank * 1.25)
    
    # Filter by source
    sub_df = df.copy()
    if src_filter == 'مرجع جامع':
        sub_df = sub_df[sub_df['منبع'].str.contains('مرجع جامع', na=False)]
    elif src_filter == 'کارنامه':
        sub_df = sub_df[sub_df['منبع'].str.contains('کارنامه', na=False)]
        
    # Filter by region and rank bounds
    filtered = sub_df[
        (sub_df['سهمیه'] == region) & 
        (sub_df['رتبه در سهمیه'] >= min_rank) & 
        (sub_df['رتبه در سهمیه'] <= max_rank)
    ].copy()
    
    if major_name != 'همه رشته‌ها':
        search_kw = major_name.replace('رادیولوژی', 'پرتوشناسی').replace('پرتو شناسی', 'پرتوشناسی')
        filtered = filtered[filtered['رشته قبولی'].str.contains(search_kw, case=False, na=False)]
        
    if filtered.empty:
        restart_markup = types.InlineKeyboardMarkup()
        restart_markup.add(types.InlineKeyboardButton("🔄 استعلام مجدد", callback_data="restart"))
        msg_text = (
            f"❌ در بازه رتبه‌ای **{min_rank:,}** الی **{max_rank:,}** (سهمیه منطقه {region})\n"
            f"قبولی ثبت‌شده‌ای برای رشته «{major_name}» در منبع منتخب یافت نشد.\n\n"
            "می‌توانید منبع را روی «همه با هم» قرار دهید یا رشته دیگری را بررسی نمایید."
        )
        bot.send_message(chat_id, msg_text, parse_mode='Markdown', reply_markup=restart_markup)
        return

    # Sort strictly by user's target major hierarchy and then by admission rank
    filtered['prio'] = filtered['رشته قبولی'].apply(get_major_priority)
    filtered.sort_values(by=['prio', 'رتبه در سهمیه'], inplace=True)
    filtered.reset_index(drop=True, inplace=True)
    filtered['ترتیب'] = range(1, len(filtered) + 1)
    
    total_count = len(filtered)
    
    header_text = (
        f"📊 **نتایج قبولی‌های پیشنهادی:**\n"
        f"👤 **سهمیه:** منطقه {region} | **رتبه داوطلب:** {rank:,}\n"
        f"🎯 **بازه تحلیلی:** ۳۰٪ خوش‌بینانه ({min_rank:,}) تا ۲۵٪ بدبینانه ({max_rank:,})\n"
        f"📂 **منبع انتخابی:** {src_label}\n"
        f"🎓 **رشته انتخابی:** {major_name}\n"
        f"📌 **تعداد کل موارد یافت‌شده:** {total_count:,} رشته‌محل\n"
        f"{'='*32}\n\n"
    )
    
    messages = []
    current_msg = header_text
    display_limit = 30
    
    for idx, (_, row) in enumerate(filtered.iterrows()):
        if idx >= display_limit:
            break
        r_val = int(row['رتبه در سهمیه'])
        if r_val < rank * 0.95:
            chance = "🎯 خوش‌بینانه (تا ۳۰٪)"
        elif r_val <= rank * 1.05:
            chance = "⚖️ محتمل و منطقی"
        else:
            chance = "🛡️ شانس بالا (حاشیه امن)"
            
        entry = (
            f"🔹 **{row['رشته قبولی']}** | {row['دوره']}\n"
            f"🏛 {row['دانشگاه قبولی']}\n"
            f"📈 آخرین رتبه: `{r_val:,}` ({chance})\n"
            f"🏷 منبع: {row['منبع']}\n"
            f"{'-'*25}\n"
        )
        if len(current_msg) + len(entry) > 3800:
            messages.append(current_msg)
            current_msg = entry
        else:
            current_msg += entry
            
    if current_msg:
        messages.append(current_msg)
        
    for msg in messages:
        bot.send_message(chat_id, msg, parse_mode='Markdown')
        
    if total_count > display_limit:
        bot.send_message(
            chat_id, 
            f"⚠️ جهت جلوگیری از شلوغی چت، {display_limit} مورد نخست در متن بالا نمایش داده شد.\n"
            f"⏳ در حال آماده‌سازی فایل اکسل شکیل و کامل شامل تمام **{total_count:,}** مورد..."
        )
        
    # Generate Styled Excel output in memory
    excel_buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "قبولی‌های پیشنهادی"
    ws.views.sheetView[0].rightToLeft = True
    
    font_name = 'B Nazanin'
    h_font = Font(name=font_name, size=11, bold=True, color='FFFFFF')
    h_fill = PatternFill(start_color='002060', end_color='002060', fill_type='solid')
    d_font = Font(name=font_name, size=11, color='000000')
    alt_fill = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
    t_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    al_center = Alignment(horizontal='center', vertical='center')
    al_right = Alignment(horizontal='right', vertical='center', indent=1)
    
    headers = ['ترتیب', 'رشته قبولی', 'دانشگاه قبولی', 'سهمیه', 'رتبه در سهمیه', 'رتبه کشوری', 'دوره', 'نیم سال', 'منبع', 'ارزیابی شانس']
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = h_font
        c.fill = h_fill
        c.alignment = al_center
        c.border = t_border
        
    for r_idx, (_, r) in enumerate(filtered.iterrows(), 2):
        r_val = int(r['رتبه در سهمیه'])
        if r_val < rank * 0.95:
            chance_text = "خوش‌بینانه (تا ۳۰٪)"
        elif r_val <= rank * 1.05:
            chance_text = "محتمل و منطقی"
        else:
            chance_text = "شانس بالا (حاشیه امن)"
            
        row_data = [
            r_idx - 1,
            r['رشته قبولی'],
            r['دانشگاه قبولی'],
            r['سهمیه'],
            r_val,
            r.get('رتبه کشوری', '-'),
            r.get('دوره', '-'),
            r.get('نیم سال', '-'),
            r.get('منبع', '-'),
            chance_text
        ]
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 22
        fill = alt_fill if r_idx % 2 == 1 else None
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = d_font
            cell.border = t_border
            if fill:
                cell.fill = fill
            h_name = headers[col_idx - 1]
            if h_name in ['ترتیب', 'سهمیه', 'رتبه در سهمیه', 'رتبه کشوری', 'دوره', 'نیم سال', 'ارزیابی شانس']:
                cell.alignment = al_center
                if isinstance(val, (int, float)) and val > 0:
                    cell.number_format = '#,##0'
            else:
                cell.alignment = al_right

    for col in ws.columns:
        max_l = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_l + 4, 12)
        
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    file_title = f"قبولی_های_رتبه_{rank}_منطقه_{region}.xlsx"
    restart_markup = types.InlineKeyboardMarkup()
    restart_markup.add(types.InlineKeyboardButton("🔄 استعلام جدید", callback_data="restart"))
    
    bot.send_document(
        chat_id, 
        excel_buffer, 
        visible_file_name=file_title, 
        caption=f"📁 **فایل اکسل مرتب و اختصاصی رتبه {rank:,} (منطقه {region})**\n\n📌 شامل تمامی گزینه‌های ممکن بر اساس ترتیب درخواستی رشته‌ها.",
        parse_mode='Markdown'
    )
    bot.send_message(chat_id, "💡 برای ارزیابی رتبه یا رشته دیگر روی دکمه زیر کلیک فرمایید:", reply_markup=restart_markup)

@bot.callback_query_handler(func=lambda call: call.data == "restart")
def callback_restart(call):
    chat_id = call.message.chat.id
    user_state[chat_id] = {}
    bot.answer_callback_query(call.id, "شروع مجدد")
    bot.send_message(
        chat_id,
        "🔄 **استعلام جدید انتخاب رشته کنکور تجربی**\n\n"
        "📍 لطفاً سهمیه منطقه خود را انتخاب فرمایید:",
        parse_mode='Markdown',
        reply_markup=get_region_keyboard()
    )


# ==============================================================================
# 🌐 KEEP-ALIVE & ANTI-SLEEP SERVER (ضد خاموشی ۲۴ ساعته برای رندر و هاست‌های ابری)
# ==============================================================================
import time
import threading
import urllib.request
from http.server import HTTPServer, BaseHTTPRequestHandler

HTML_STATUS_PAGE = """<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>ربات انتخاب رشته کنکور - وضعیت ۲۴ ساعته</title>
    <style>
        body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #0f172a; color: #f8fafc; text-align: center; padding-top: 60px; margin: 0; }
        .card { background: #1e293b; max-width: 520px; margin: 0 auto; padding: 30px; border-radius: 16px; box-shadow: 0 10px 25px rgba(0,0,0,0.5); border: 1px solid #334155; }
        .status-badge { display: inline-block; background: #10b981; color: white; padding: 6px 16px; border-radius: 20px; font-weight: bold; margin-bottom: 20px; font-size: 14px; }
        h1 { color: #38bdf8; font-size: 22px; margin-bottom: 12px; }
        p { color: #94a3b8; font-size: 14px; line-height: 1.6; }
        .tag { font-family: monospace; background: #334155; padding: 2px 6px; border-radius: 4px; color: #f1f5f9; }
    </style>
</head>
<body>
    <div class="card">
        <div class="status-badge">● ۲۴ ساعته آنلاین و فعال (Active 24/7)</div>
        <h1>ربات تلگرام انتخاب رشته کنکور</h1>
        <p>وب‌سرور پایدارساز و سیستم ضد خاموشی خودکار (<span class="tag">Keep-Alive</span>) با موفقیت در حال اجرا است.</p>
        <p style="margin-top: 15px; font-size: 12px; color: #64748b;">Render Sleep Preventer & Health Monitor Active</p>
    </div>
</body>
</html>"""

class KeepAliveHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-type", "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(HTML_STATUS_PAGE.strip().encode("utf-8"))

    def do_HEAD(self):
        self.send_response(200)
        self.send_header("Content-type", "text/html; charset=utf-8")
        self.end_headers()

    def log_message(self, format, *args):
        pass

def run_keep_alive_server(port):
    try:
        server_address = ("0.0.0.0", port)
        httpd = HTTPServer(server_address, KeepAliveHandler)
        print(f"🌐 [Keep-Alive] وب‌سرور پایش روی پورت {port} فعال شد.")
        httpd.serve_forever()
    except Exception as e:
        print(f"⚠️ [Keep-Alive] وب‌سرور: {e}")

def auto_pinger_loop():
    time.sleep(30)
    while True:
        url = os.environ.get("RENDER_EXTERNAL_URL") or os.environ.get("APP_URL")
        now_str = time.strftime("%H:%M:%S")
        if url:
            try:
                if not url.startswith("http"):
                    url = "https://" + url
                req = urllib.request.Request(
                    url,
                    headers={"User-Agent": "Mozilla/5.0 (RenderKeepAlive/1.0)"}
                )
                with urllib.request.urlopen(req, timeout=20) as response:
                    if response.status == 200:
                        print(f"💓 [Keep-Alive] پینگ خودکار موفق به {url} در {now_str}")
            except Exception as e:
                print(f"⚠️ [Keep-Alive] وضعیت پینگ: {e}")
        time.sleep(600)

def start_keep_alive():
    port = int(os.environ.get("PORT", 10000))
    t1 = threading.Thread(target=run_keep_alive_server, args=(port,), daemon=True)
    t1.start()
    t2 = threading.Thread(target=auto_pinger_loop, daemon=True)
    t2.start()


if __name__ == '__main__':
    try:
        start_keep_alive()
    except Exception as e:
        print(f"⚠️ خطا در راه‌اندازی Keep-Alive: {e}")
    print("🚀 ربات با موفقیت و بر اساس دکمه‌های اینلاین آماده اجرا شد!")
    bot.infinity_polling()
